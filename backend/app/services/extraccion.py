"""Extracción de texto de los archivos de evidencia (PDF, DOCX, PPTX, texto, imágenes)."""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from html import unescape

MAX_ADJUNTO_BYTES = 20 * 1024 * 1024   # límite razonable para enviar un PDF/imagen al modelo
IMAGENES = {"image/png", "image/jpeg", "image/webp", "image/gif"}
EXT_MIME = {
    ".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".txt": "text/plain", ".md": "text/markdown", ".csv": "text/csv", ".json": "application/json",
    ".html": "text/html", ".htm": "text/html",
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif",
}


@dataclass
class Adjunto:
    """Archivo que se envía tal cual al modelo (imagen o PDF escaneado)."""
    kind: str            # "image" | "document"
    mime: str
    data: bytes
    nombre: str


@dataclass
class Extraccion:
    texto: str = ""
    adjunto: Adjunto | None = None
    aviso: str | None = None


def _mime(nombre: str, mime: str | None) -> str:
    mime = (mime or "").split(";")[0].strip().lower()
    if mime and mime != "application/octet-stream":
        return mime
    m = re.search(r"\.[A-Za-z0-9]+$", nombre or "")
    return EXT_MIME.get(m.group(0).lower(), mime) if m else mime


def _limpiar(texto: str) -> str:
    texto = texto.replace("\x00", "")
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()


def _pdf(data: bytes, nombre: str) -> Extraccion:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        paginas = [(p.extract_text() or "") for p in reader.pages]
    except Exception as e:  # PDF corrupto o cifrado
        return Extraccion(aviso=f"No se pudo leer el PDF «{nombre}»: {e.__class__.__name__}")
    texto = _limpiar("\n\n".join(paginas))
    if len(texto) >= 80:
        return Extraccion(texto=texto)
    # Poco o nada de texto: probablemente escaneado. Se envía el PDF al modelo (lee imágenes).
    if len(data) <= MAX_ADJUNTO_BYTES:
        return Extraccion(
            texto=texto,
            adjunto=Adjunto("document", "application/pdf", data, nombre),
            aviso=f"«{nombre}» parece un PDF escaneado; se enviará como documento al modelo.",
        )
    return Extraccion(aviso=f"«{nombre}» es un PDF escaneado demasiado grande; se omitió.")


def _docx(data: bytes, nombre: str) -> Extraccion:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        return Extraccion(aviso=f"No se pudo leer el Word «{nombre}»: {e.__class__.__name__}")
    partes = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for fila in t.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))
    return Extraccion(texto=_limpiar("\n".join(partes)))


def _pptx(data: bytes, nombre: str) -> Extraccion:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            slides = sorted(
                (n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                key=lambda n: int(re.search(r"(\d+)\.xml$", n).group(1)),
            )
            partes = []
            for i, n in enumerate(slides, 1):
                xml = z.read(n).decode("utf-8", "replace")
                textos = [unescape(t) for t in re.findall(r"<a:t>([^<]*)</a:t>", xml)]
                if textos:
                    partes.append(f"[Diapositiva {i}] " + " ".join(textos))
    except Exception as e:
        return Extraccion(aviso=f"No se pudo leer la presentación «{nombre}»: {e.__class__.__name__}")
    return Extraccion(texto=_limpiar("\n".join(partes)))


def _texto_plano(data: bytes, html: bool = False) -> Extraccion:
    for enc in ("utf-8", "latin-1"):
        try:
            s = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        s = data.decode("utf-8", "replace")
    if html:
        s = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", s)
        s = unescape(re.sub(r"(?s)<[^>]+>", " ", s))
    return Extraccion(texto=_limpiar(s))


def extraer(nombre: str, mime: str | None, data: bytes) -> Extraccion:
    """Devuelve el texto (y/o un adjunto) de un archivo. Nunca lanza: informa con `aviso`."""
    nombre = nombre or "archivo"
    tipo = _mime(nombre, mime)

    if tipo == "application/pdf":
        return _pdf(data, nombre)
    if tipo.endswith("wordprocessingml.document"):
        return _docx(data, nombre)
    if tipo.endswith("presentationml.presentation"):
        return _pptx(data, nombre)
    if tipo in IMAGENES:
        if len(data) > MAX_ADJUNTO_BYTES:
            return Extraccion(aviso=f"La imagen «{nombre}» es demasiado grande; se omitió.")
        return Extraccion(adjunto=Adjunto("image", tipo, data, nombre))
    if tipo in ("text/html", "application/xhtml+xml"):
        return _texto_plano(data, html=True)
    if tipo.startswith("text/") or tipo in ("application/json", "application/xml"):
        return _texto_plano(data)
    return Extraccion(aviso=f"Tipo de archivo no soportado para generar preguntas: «{nombre}» ({tipo or 'desconocido'}).")
